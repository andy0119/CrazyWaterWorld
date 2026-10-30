import sys
import unittest
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'backend'))
from gift_excel import parse_excel, export_excel, MAX_BYTES
from openpyxl import Workbook, load_workbook


def sample(rows, title='礼包码'):
    book=Workbook();book.active.title=title
    for row in rows:book.active.append(row)
    out=BytesIO();book.save(out);book.close();return out.getvalue()


class GiftExcel(unittest.TestCase):
    def test_roundtrip_preserves_identifiers(self):
        codes=['001234','AbC-123','abc-123','=LITERAL','+LITERAL','@LITERAL']
        self.assertEqual(codes,parse_excel(export_excel(codes),'codes.xlsx')['codes'])
        book=load_workbook(BytesIO(export_excel(codes)),read_only=True,data_only=False)
        self.assertTrue(all(c.data_type=='s' for row in book.active.iter_rows(min_row=2) for c in row if c.value is not None));book.close()

    def test_template_and_export_share_single_column_format(self):
        root=Path(__file__).resolve().parents[1]
        template=load_workbook(root/'web/gift-codes-template.xlsx')
        exported=load_workbook(BytesIO(export_excel(['ABC','00123'])))
        a,b=template.active,exported.active
        self.assertEqual(a.max_column,1);self.assertEqual(b.max_column,1)
        self.assertEqual(a['A1'].value,b['A1'].value)
        self.assertEqual(a.column_dimensions['A'].width,b.column_dimensions['A'].width)
        self.assertEqual(a.freeze_panes,b.freeze_panes)
        self.assertEqual(a.sheet_view.showGridLines,b.sheet_view.showGridLines)
        self.assertEqual(a.auto_filter.ref,b.auto_filter.ref)
        for row in [1,2,3]:
            x,y=a.cell(row,1),b.cell(row,1)
            self.assertEqual(x.font.name,y.font.name);self.assertEqual(x.font.sz,y.font.sz)
            self.assertEqual(bool(x.font.bold),bool(y.font.bold))
            self.assertEqual(x.number_format,y.number_format)
            self.assertEqual(x.alignment.horizontal,y.alignment.horizontal)
            self.assertEqual(x.alignment.vertical,y.alignment.vertical)
            self.assertEqual(a.row_dimensions[row].height,b.row_dimensions[row].height)
        self.assertEqual(a['A1'].fill.fgColor.rgb[-6:],b['A1'].fill.fgColor.rgb[-6:])
        template.close();exported.close()

    def test_expired_archive_is_separate_and_not_reimported(self):
        data=export_excel(['LIVE','DEAD'],['DEAD','EXPIRED2'])
        book=load_workbook(BytesIO(data),read_only=True)
        self.assertEqual(book.sheetnames,['礼包码','過期'])
        self.assertEqual([c.value for row in book['礼包码'].iter_rows(min_row=2) for c in row if c.value],['LIVE'])
        self.assertEqual([c.value for row in book['過期'].iter_rows(min_row=2) for c in row if c.value],['DEAD','EXPIRED2'])
        book.close()
        self.assertEqual(parse_excel(data,'archive.xlsx')['codes'],['LIVE'])
        with self.assertRaises(ValueError):parse_excel(sample([['礼包码'],['DEAD']],title='過期'),'archive.xlsx')

    def test_dedupe_blank_header_and_examples(self):
        result=parse_excel(sample([['禮包碼','备注'],[' EXAMPLE_ONLY_001 '],['AbC'],['AbC'],[None],['00123']]),'list.xlsx')
        self.assertEqual(result['codes'],['AbC','00123']);self.assertEqual(result['duplicates'],1);self.assertEqual(result['examples_skipped'],1)

    def test_headerless_and_moved_code_column(self):
        self.assertEqual(parse_excel(sample([['ABC'],['DEF']]),'file.xlsx')['codes'],['ABC','DEF'])
        self.assertEqual(parse_excel(sample([['备注','code'],['note','ABC']]),'file.xlsx')['codes'],['ABC'])
        self.assertEqual(parse_excel(sample([[None,'code'],[None,'ABC']]),'file.xlsx')['codes'],['ABC'])

    def test_duplicate_rows_do_not_use_distinct_code_limit(self):
        self.assertEqual(parse_excel(sample([['礼包码']]+[['SAME']]*1001),'file.xlsx')['codes'],['SAME'])

    def test_rejects_formulas_numeric_and_oversized_lists(self):
        for rows in [[['礼包码'],['=1+1']],[['礼包码'],[12345]],[['礼包码'],['x'*101]]]:
            with self.assertRaises(ValueError):parse_excel(sample(rows),'codes.xlsx')
        self.assertEqual(len(parse_excel(export_excel([str(i) for i in range(1500)]),'large.xlsx')['codes']),1500)

    def test_invalid_empty_and_unsupported_files(self):
        for data,name in [(b'not an xlsx','a.xlsx'),(b'','a.xlsx'),(b'x','a.xls'),(b'x'*(MAX_BYTES+1),'a.xlsx'),(sample([['礼包码']]),'a.xlsx')]:
            with self.assertRaises(ValueError):parse_excel(data,name)

    def test_rejects_large_inflated_archive_and_xml_entities(self):
        for name,content in [('xl/fake.xml',b'x'*(11*1024*1024)),('xl/fake.xml',b'<!DOCTYPE workbook [<!ENTITY a "x">]>')]:
            stream=BytesIO()
            with ZipFile(stream,'w',ZIP_DEFLATED) as z:z.writestr(name,content)
            with self.assertRaises(ValueError):parse_excel(stream.getvalue(),'a.xlsx')

    def test_rejects_sparse_rows_instead_of_silently_truncating(self):
        book=Workbook();book.active['A1']='礼包码';book.active['A10003']='LATE';out=BytesIO();book.save(out);book.close()
        with self.assertRaises(ValueError):parse_excel(out.getvalue(),'a.xlsx')


if __name__=='__main__':unittest.main()
