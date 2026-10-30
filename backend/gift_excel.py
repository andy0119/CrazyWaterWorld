"""Excel exchange for candidate gift codes; never redeems codes or changes history."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import re

MAX_BYTES = 2 * 1024 * 1024
EXAMPLES = {'EXAMPLE_ONLY_001', 'EXAMPLE_ONLY_002'}
HEADERS = {'礼包码', '禮包碼', '兑换码', '兌換碼', 'code', 'giftcode', 'gift_code'}


def clean_codes(values):
    if not isinstance(values, list):
        raise ValueError('礼包码须为文字清单')
    result = [];seen=set()
    for value in values:
        if not isinstance(value, str) or not 0 < len(value.strip()) <= 100 or re.search(r'[\x00-\x1f]', value):
            raise ValueError('每组礼包码须为 1 至 100 字的文字，不能包含换行或控制字符')
        value = value.strip()
        if value not in seen:
            result.append(value);seen.add(value)
    return result


def parse_excel(data, filename):
    if Path(filename).suffix.lower() != '.xlsx':
        raise ValueError('请使用 .xlsx 文件；旧版 .xls 请先在 Excel 另存为 .xlsx')
    if not data or len(data) > MAX_BYTES:
        raise ValueError('Excel 文件不可为空，且不能超过 2 MB')
    try:
        with ZipFile(BytesIO(data)) as archive:
            files = archive.infolist()
            if len(files) > 200 or sum(f.file_size for f in files) > 10 * 1024 * 1024:
                raise ValueError('Excel 内容过大，请只保留礼包码清单')
            for entry in files:
                if entry.flag_bits & 1 or entry.filename.lower().endswith('vbaproject.bin'):
                    raise ValueError('不支持加密或含宏的文件')
                if entry.filename.lower().endswith('.xml'):
                    xml = archive.read(entry).upper()
                    if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml:
                        raise ValueError('Excel XML 格式不支持')
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('无法读取 Excel，请确认文件未损坏且格式为 .xlsx') from exc
    try:
        if len(book.worksheets) > 10:
            raise ValueError('工作表过多，请只保留礼包码清单')
        sheet = next((s for s in book.worksheets if s.title in ('礼包码', '禮包碼')), book.worksheets[0])
        if sheet.title in ('過期','过期'):
            raise ValueError('只有过期工作表，未导入；请使用礼包码主表')
        if sheet.max_row and sheet.max_row > 10002:
            raise ValueError('工作表行数过多，请只保留礼包码清单（最多 10000 行）')
        # Bound iteration independently of the sheet's declared dimension.
        first = next(sheet.iter_rows(min_row=1, max_row=1, max_col=20), ())
        column = next((c.column for c in first if isinstance(c.value, str) and c.value.strip().lower() in HEADERS), 1)
        header = first[column-1] if len(first) >= column else None
        start = 2 if header and isinstance(header.value, str) and header.value.strip().lower() in HEADERS else 1
        values, errors, examples = [], [], 0
        for cells in sheet.iter_rows(min_row=start, max_row=10002, min_col=column, max_col=column):
            cell = cells[0]
            value = cell.value
            if value is None or isinstance(value, str) and not value.strip():
                continue
            if cell.data_type == 'f':
                errors.append(f'第 {cell.row} 行含公式，请贴上文字值')
            elif not isinstance(value, str):
                errors.append(f'第 {cell.row} 行不是文字，请把礼包码栏设为文字以保留开头的 0')
            elif value.strip() in EXAMPLES:
                examples += 1
            else:
                values.append(value)
            if len(errors) >= 5:
                break
            if len(values) > 10000:
                raise ValueError('文件最多可包含 10000 行')
        if errors:
            raise ValueError('；'.join(errors))
        codes = clean_codes(values)
        if not codes:
            raise ValueError('没有可导入的礼包码；请先把范例占位码换成真实礼包码')
        return {'codes': codes, 'duplicates': len(values) - len(codes), 'examples_skipped': examples, 'sheet': sheet.title}
    finally:
        book.close()


def export_excel(values, expired=()):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    expired = clean_codes(list(expired))
    excluded = set(expired)
    codes = [c for c in clean_codes(values) if c not in excluded]
    book = Workbook()
    for index, (title, entries) in enumerate([('礼包码', codes), ('過期', expired)]):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = title
        sheet.append(['礼包码'])
        sheet.column_dimensions['A'].width = 42
        sheet.freeze_panes = 'A2'
        sheet.sheet_view.showGridLines = False
        sheet['A1'].font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        sheet['A1'].fill = PatternFill('solid', fgColor='24466B')
        sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')
        sheet.row_dimensions[1].height = 22
        for row in range(2, max(1002, len(entries)+2)):
            cell = sheet.cell(row, 1)
            cell.number_format = '@'
            cell.font = Font(name='Arial', size=11)
            cell.alignment = Alignment(horizontal='left', vertical='center')
            sheet.row_dimensions[row].height = 22
        for row, code in enumerate(entries, 2):
            cell = sheet.cell(row, 1, code)
            cell.data_type = 's'
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def mount(app, auth, conn=None, own_sub=None):
    from fastapi import Header, HTTPException, Request, Response

    @app.post('/api/gift-codes/import-excel')
    async def import_file(request: Request, filename: str = 'codes.xlsx', authorization: str | None = Header(None)):
        auth(authorization)
        chunks, size = [], 0
        async for chunk in request.stream():
            size += len(chunk)
            if size > MAX_BYTES:
                raise HTTPException(413, 'Excel 文件不能超过 2 MB')
            chunks.append(chunk)
        try:
            return parse_excel(b''.join(chunks), filename)
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc

    @app.post('/api/gift-codes/export-excel')
    async def export_file(request: Request, authorization: str | None = Header(None)):
        auth(authorization)
        chunks, size = [], 0
        async for chunk in request.stream():
            size += len(chunk)
            if size > 512 * 1024:
                raise HTTPException(413, '礼包码清单过大')
            chunks.append(chunk)
        try:
            import json
            payload = json.loads(b''.join(chunks))
            codes = payload.get('codes') if isinstance(payload, dict) else None
            expired = []
            if conn is not None:
                account = payload.get('account') if isinstance(payload, dict) else None
                if type(account) is not int or account <= 0:
                    raise ValueError('请从当前账号的礼包码设置导出')
                own_sub(account, authorization)
                from gift_history import expired_codes
                with conn() as db:expired = sorted(expired_codes(db, account))
            data = export_excel(codes, expired)
        except (ValueError, TypeError) as exc:
            raise HTTPException(422, str(exc)) from exc
        return Response(data, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={
            'Content-Disposition': 'attachment; filename="gift-codes.xlsx"', 'Cache-Control': 'no-store'})
